"""Read the supplied review without changing the originals."""
from pathlib import Path
import json
from pypdf import PdfReader
from pptx import Presentation

SOURCE = Path('C:/Users/tae06/OneDrive - postech.ac.kr/PSI - 02 2026 UGRP 3팀')
OUT = Path(__file__).parent
doc = PdfReader(SOURCE / '이태호 literature review.pdf')
slides = Presentation(SOURCE / '이태호 literature review.pptx')
records = []
for i, page in enumerate(doc.pages, 1):
    records.append({'page': i, 'pdf_text': page.extract_text()})
for i, slide in enumerate(slides.slides, 1):
    record = records[i-1] if i <= len(records) else {'page': i}
    record['pptx_text'] = '\n'.join(s.text for s in slide.shapes if s.has_text_frame)
    record['notes'] = slide.notes_slide.notes_text_frame.text if slide.has_notes_slide else ''
    if i > len(records):
        records.append(record)
(OUT / 'source_review_extraction.json').write_text(json.dumps(records, ensure_ascii=False, indent=2), encoding='utf-8')
for record in records:
    print(f"\n--- PAGE {record['page']} ---\n{record.get('pdf_text', '')}\nNOTES: {record['notes']}")
print(f'PDF pages: {len(doc.pages)}; PPTX slides: {len(slides.slides)}')

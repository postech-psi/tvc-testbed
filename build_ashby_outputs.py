from pathlib import Path
from collections import Counter
import textwrap

import openpyxl
from PIL import Image, ImageDraw, ImageFont
import random


SOURCE = Path(r"C:\Users\tae06\Downloads\TVC_Ashby_Research_Corpus.xlsx")
OUT = Path(r"C:\Users\tae06\CODE\tvc-testbed\research_outputs")
OUT.mkdir(parents=True, exist_ok=True)


def clean(value, fallback="Not reported"):
    if value is None or str(value).strip() == "":
        return fallback
    return str(value).strip()


def esc(value):
    return clean(value).replace("|", "\\|").replace("\n", " ")


wb = openpyxl.load_workbook(SOURCE, data_only=True)
ws = wb["Papers"]
headers = [cell.value for cell in ws[4]]
rows = []
for excel_row in ws.iter_rows(min_row=5, max_row=49, values_only=True):
    row = dict(zip(headers, excel_row))
    if row.get("ID"):
        rows.append(row)

gap_ws = wb["Gap Matrix"]
gaps = []
for r in gap_ws.iter_rows(min_row=6, values_only=True):
    if r[0]:
        gaps.append({"region": r[0], "assessment": r[1], "evidence": r[2], "implication": r[3]})

# Condense closely related control labels only for legible plot color encoding.
def control_group(name):
    s = clean(name, "Other").lower()
    if "reinforcement" in s or "rl-" in s or "safe learning" in s or "residual" in s:
        return "Learning / RL"
    if "optimization" in s or "mpc" in s or "optimal guidance" in s:
        return "Optimization / MPC"
    if "adaptive" in s or "robust" in s or "nonlinear" in s or "geometric" in s:
        return "Robust / adaptive / nonlinear"
    if "identification" in s or "benchmark" in s or "review" in s:
        return "Identification / review"
    return "Classical / model-based"


ring_markers = {"1 Exact": "o", "2 Close": "s", "3 Analogous": "^", "4 Method": "D"}
colors = {
    "Learning / RL": "#d95f02",
    "Optimization / MPC": "#1b9e77",
    "Robust / adaptive / nonlinear": "#7570b3",
    "Identification / review": "#666666",
    "Classical / model-based": "#1f78b4",
}
scope_labels = {1: "Attitude", 2: "Pose", 3: "Trajectory", 4: "Online guidance", 5: "Integrated mission"}
validation_labels = {0: "Review", 1: "Simulation", 2: "SIL / HIL", 3: "Constrained rig", 4: "Indoor free flight", 5: "Outdoor free flight", 6: "Operational / on-orbit"}
authority_labels = {0: "None", 1: "Learned model / adaptation", 2: "Residual / high-level / safety", 3: "Direct / integrated policy"}

W, H = 2400, 1320
img = Image.new("RGB", (W, H), "#f7f5ef")
draw = ImageDraw.Draw(img)
font_path = r"C:\Windows\Fonts\arial.ttf"
bold_path = r"C:\Windows\Fonts\arialbd.ttf"
font = ImageFont.truetype(font_path, 18)
small = ImageFont.truetype(font_path, 15)
tiny = ImageFont.truetype(font_path, 13)
title_font = ImageFont.truetype(bold_path, 34)
panel_font = ImageFont.truetype(bold_path, 23)

draw.text((55, 35), "Thrust-Vector Control Research Landscape", fill="#202020", font=title_font)
draw.text((55, 85), "45 evidence-coded papers. Labels are workbook IDs; jitter only separates overlapping categories.", fill="#555555", font=font)

panels = [(90, 175, 1110, 1005), (1290, 175, 2310, 1005)]

def marker(cx, cy, shape, fill, radius=8):
    edge = "#202020"
    if shape == "o":
        draw.ellipse((cx-radius, cy-radius, cx+radius, cy+radius), fill=fill, outline=edge, width=2)
    elif shape == "s":
        draw.rectangle((cx-radius, cy-radius, cx+radius, cy+radius), fill=fill, outline=edge, width=2)
    elif shape == "^":
        draw.polygon([(cx, cy-radius-2), (cx-radius-2, cy+radius), (cx+radius+2, cy+radius)], fill=fill, outline=edge)
    else:
        draw.polygon([(cx, cy-radius-2), (cx-radius-2, cy), (cx, cy+radius+2), (cx+radius+2, cy)], fill=fill, outline=edge)

def draw_panel(box, xcodes, xlabels, title, key, seed, highlight=False):
    left, top, right, bottom = box
    draw.rectangle(box, fill="#fffdfa")
    px0, py0, px1, py1 = left+170, top+75, right-35, bottom-150
    if highlight:
        hx0 = px0 + (3-min(xcodes))/(max(xcodes)-min(xcodes))*(px1-px0) - 75
        hx1 = hx0 + 150
        hy_top = py1 - 5/6*(py1-py0) - 45
        hy_bottom = py1 - 4/6*(py1-py0) + 45
        draw.rectangle((hx0, hy_top, hx1, hy_bottom), fill="#fff1c7")
        draw.multiline_text((hx0-15, hy_top-72), "Sparse target zone:\ndirect policy + free flight", fill="#6b4f00", font=tiny, align="center")
    for y in range(7):
        py = py1 - y/6*(py1-py0)
        draw.line((px0, py, px1, py), fill="#dedbd2", width=1)
        label = validation_labels[y]
        tw = draw.textbbox((0,0), label, font=small)[2]
        draw.text((px0-tw-14, py-9), label, fill="#333333", font=small)
    for i, x in enumerate(xcodes):
        px = px0 + i/(len(xcodes)-1)*(px1-px0)
        draw.line((px, py0, px, py1), fill="#e7e3da", width=1)
        label = xlabels[x]
        wrapped = "\n".join(textwrap.wrap(label, 16))
        bbox = draw.multiline_textbbox((0,0), wrapped, font=tiny, align="center")
        draw.multiline_text((px-(bbox[2]-bbox[0])/2, py1+15), wrapped, fill="#333333", font=tiny, align="center", spacing=3)
    draw.line((px0, py0, px0, py1), fill="#444444", width=2)
    draw.line((px0, py1, px1, py1), fill="#444444", width=2)
    draw.text((left+15, top+15), title, fill="#202020", font=panel_font)
    rr = random.Random(seed)
    for row in rows:
        xv = int(row[key]); yv = int(row["Validation Code"])
        xi = xcodes.index(xv)
        px = px0 + xi/(len(xcodes)-1)*(px1-px0) + rr.uniform(-24,24)
        py = py1 - yv/6*(py1-py0) + rr.uniform(-16,16)
        marker(px, py, ring_markers[row["Relevance Ring"]], colors[control_group(row["Control Family"])])
        draw.text((px+10, py-17), row["ID"], fill="#202020", font=tiny)

draw_panel(panels[0], list(scope_labels), scope_labels, "A. Functional breadth vs. validation evidence", "Functional Scope Code", 13)
draw_panel(panels[1], list(authority_labels), authority_labels, "B. Learning authority vs. validation evidence", "Learning Authority Code", 29, True)

legend_y = 1090
draw.text((75, legend_y), "Control family (color)", fill="#202020", font=ImageFont.truetype(bold_path, 16))
x = 75
for label, color in colors.items():
    marker(x+8, legend_y+42, "o", color, 7)
    draw.text((x+23, legend_y+31), label, fill="#333333", font=tiny)
    x += 255 if len(label) < 23 else 315
draw.text((75, legend_y+82), "Architectural relevance (shape)", fill="#202020", font=ImageFont.truetype(bold_path, 16))
x = 75
for label, shape in ring_markers.items():
    marker(x+8, legend_y+123, shape, "#dddddd", 7)
    draw.text((x+23, legend_y+112), label, fill="#333333", font=tiny)
    x += 190
draw.text((1430, legend_y+105), "Highlighted zone is a search target, not proof of universal absence.", fill="#6b4f00", font=small)

png_path = OUT / "tvc_ashby_landscape.png"
img.save(png_path, dpi=(180, 180))

# Markdown catalog.
ring_counts = Counter(row["Relevance Ring"] for row in rows)
validation_counts = Counter(int(row["Validation Code"]) for row in rows)
lines = [
    "# TVC Ashby research reader",
    "",
    f"Source workbook: `{SOURCE}`",
    "",
    f"This reader indexes **{len(rows)} papers** using the same IDs as the Ashby plot. It is designed for reading the corpus one paper at a time, without treating blank specifications as zero or treating method-transfer papers as direct competitors.",
    "",
    "![TVC Ashby landscape](tvc_ashby_landscape.png)",
    "",
    "## How to read the plot",
    "",
    "- **Panel A** maps demonstrated functional breadth against validation maturity.",
    "- **Panel B** maps how much authority learning has in the control loop against validation maturity.",
    "- **Color** is a plot-only grouping of related control families; the original detailed family is preserved in every paper entry.",
    "- **Marker shape** shows architectural relevance: exact, close, analogous, or method transfer.",
    "- Point locations are categorical. Small deterministic jitter only separates overlaps; it has no quantitative meaning.",
    "- The highlighted region is a search target, not proof that no work exists outside this corpus.",
    "",
    "## Coding scales",
    "",
    "| Code | Functional scope | Validation | Learning authority |",
    "|---:|---|---|---|",
]
for code in range(7):
    lines.append(f"| {code} | {scope_labels.get(code, '—')} | {validation_labels.get(code, '—')} | {authority_labels.get(code, '—')} |")

lines += [
    "",
    "## Corpus at a glance",
    "",
    f"- Relevance rings: " + ", ".join(f"{k}: {ring_counts[k]}" for k in ring_markers),
    f"- Free-flight or operational evidence (codes 4–6): **{sum(validation_counts[c] for c in (4,5,6))} papers**.",
    f"- Simulation-only evidence (code 1): **{validation_counts[1]} papers**.",
    "- This is an evidence map, not a ranking of algorithm quality.",
    "",
    "## Evidence-based gap matrix",
    "",
    "The statements below are claims about this corpus only.",
    "",
    "| Candidate region | Assessment | Evidence | Implication |",
    "|---|---|---|---|",
]
for gap in gaps:
    lines.append(f"| {esc(gap['region'])} | {esc(gap['assessment'])} | {esc(gap['evidence'])} | {esc(gap['implication'])} |")

lines += ["", "## Paper-by-paper catalog", ""]

for ring in ring_markers:
    ring_rows = [r for r in rows if r["Relevance Ring"] == ring]
    lines += [f"## {ring} ({len(ring_rows)} papers)", ""]
    for r in ring_rows:
        source_url = clean(r["Primary Source URL"], "")
        title = clean(r["Paper Title"])
        if source_url:
            title_line = f"### {r['ID']} — [{title}]({source_url})"
        else:
            title_line = f"### {r['ID']} — {title}"
        lines += [
            title_line,
            "",
            f"**Citation:** {clean(r['First Author / Citation'])} ({clean(r['Year'])}), {clean(r['Venue'])}.  ",
            f"**Architecture:** {clean(r['Architecture Family'])}; {clean(r['Vehicle Form'])}.  ",
            f"**Actuation:** {clean(r['Propulsion'])}; {clean(r['Vectoring Topology'])}; {clean(r['Actuation Class'])}.  ",
            f"**Control:** {clean(r['Control Family'])} — {clean(r['Control Method'])}.  ",
            f"**Learning:** {clean(r['Learning Role'])}; algorithm: {clean(r['RL Algorithm'])}; authority code: {clean(r['Learning Authority Code'])}.  ",
            f"**GNC scope:** {clean(r['Guidance'])}; estimation: {clean(r['Estimation'])}; function: {clean(r['Functional Scope'])} (code {clean(r['Functional Scope Code'])}).  ",
            f"**Validation:** {clean(r['Validation Level'])} (code {clean(r['Validation Code'])}); free flight: {clean(r['Free Flight'])}; environment: {clean(r['Environment'])}.  ",
            f"**Constraints and uncertainty:** {clean(r['Constraint Handling'])}; {clean(r['Uncertainty Treatment'])}.  ",
            f"**Identification / compute:** {clean(r['System Identification'])}; onboard compute: {clean(r['Onboard Compute'])}; control rate: {clean(r['Control Rate (Hz)'])}.  ",
            f"**Reported evidence:** {clean(r['Key Metrics'])}; hardware: {clean(r['Hardware Evidence'])}; open source: {clean(r['Open Source'])}.  ",
            f"**Why it matters:** {clean(r['Project Relevance'])}  ",
            f"**Evidence status:** {clean(r['Evidence Status'])}. **Uncertainty:** {clean(r['Uncertainty Flag'])}.  ",
            f"**Corpus note:** {clean(r['Evidence Note'])}  ",
            f"**Identifier:** {clean(r['DOI / arXiv'])}.",
            "",
        ]

lines += [
    "## Suggested reading sequence",
    "",
    "1. Read the **Exact** ring first to understand the current physical design space and its experimental baseline.",
    "2. Read the **Close** ring next for alternate thrust-vectoring architectures and allocation methods.",
    "3. Use the **Analogous** ring for transferable landing, marine, and aerial GNC evidence.",
    "4. Use the **Method** ring only after identifying a specific mechanism to transfer, such as residual RL, safe learning, active identification, or sim-to-real validation.",
    "",
    "## Important limitations",
    "",
    "- Most entries are abstract-verified rather than full-text verified.",
    "- Missing numerical specifications remain `Not reported`; they were not inferred.",
    "- The plot does not establish novelty by itself. Any sparse region should be checked through backward and forward citation searches before making a publication claim.",
]

md_path = OUT / "TVC_Ashby_Research_Reader.md"
md_path.write_text("\n".join(lines), encoding="utf-8")

print(f"Created {png_path}")
print(f"Created {md_path}")
print(f"Markdown lines: {len(lines)}")

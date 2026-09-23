from pathlib import Path
from collections import Counter
import json
import random
import textwrap

import openpyxl
from PIL import Image, ImageDraw, ImageFont

SRC = Path(r"C:\Users\tae06\Downloads\TVC_Ashby_Research_Corpus.xlsx")
OUT = Path(r"C:\Users\tae06\CODE\tvc-testbed\research_outputs\curated")
OUT.mkdir(parents=True, exist_ok=True)

wb = openpyxl.load_workbook(SRC, data_only=True)
ws = wb["Papers"]
headers = [c.value for c in ws[4]]
original = [dict(zip(headers, r)) for r in ws.iter_rows(min_row=5, max_row=49, values_only=True)]
original = [r for r in original if r.get("ID")]
by_id = {r["ID"]: r for r in original}

selected_ids = ["R01", "R02", "R04", "R05", "R06", "R09", "R10", "R13", "R20"]

def compact_existing(r):
    pub_url = r["Primary Source URL"]
    doi = str(r.get("DOI / arXiv") or "")
    if doi.startswith("10."):
        pub_url = "https://doi.org/" + doi
    return {
        "ID": r["ID"], "Year": r["Year"], "Title": r["Paper Title"],
        "Citation": r["First Author / Citation"], "Venue": r["Venue"],
        "Publication Type": r["Publication Type"], "Peer Reviewed": r["Peer Reviewed"],
        "Architecture": r["Architecture Family"], "Vehicle Form": r["Vehicle Form"],
        "Vectoring": r["Vectoring Topology"], "Control Family": r["Control Family"],
        "Control Method": r["Control Method"], "Learning Role": r["Learning Role"],
        "Functional Scope": r["Functional Scope"], "Functional Scope Code": int(r["Functional Scope Code"]),
        "Validation": r["Validation Level"], "Validation Code": int(r["Validation Code"]),
        "Hardware Evidence": r["Hardware Evidence"], "Free Flight": r["Free Flight"],
        "Key Metrics": r["Key Metrics"], "Project Relevance": r["Project Relevance"],
        "Evidence Status": r["Evidence Status"], "Evidence Note": r["Evidence Note"],
        "DOI": doi, "Source URL": pub_url,
    }

main = [compact_existing(by_id[i]) for i in selected_ids]
for r in main:
    if r["ID"] == "R10":
        r.update({"Venue":"IEEE Robotics and Automation Letters","Publication Type":"Journal","DOI":"10.1109/LRA.2024.3451391","Source URL":"https://doi.org/10.1109/LRA.2024.3451391"})
    if r["ID"] == "R20":
        r.update({"Citation":"Xie, Xian and Gu","DOI":"10.1016/j.isatra.2022.06.006","Source URL":"https://doi.org/10.1016/j.isatra.2022.06.006","Evidence Status":"Publisher record verified"})

new_records = [
    {
        "ID":"N01","Year":2024,"Title":"Development of a Tube-Launched Tail-Sitter Unmanned Aerial Vehicle",
        "Citation":"Cai et al.","Venue":"International Journal of Micro Air Vehicles","Publication Type":"Journal","Peer Reviewed":"Yes",
        "Architecture":"Single coaxial propulsion unit on a two-axis gimbal; tail-sitter airframe","Vehicle Form":"Tube-launched coaxial tail-sitter",
        "Vectoring":"Two-axis gimbal, ±30°; differential rotor RPM for yaw","Control Family":"Classical",
        "Control Method":"Cascaded attitude/rate PID on custom autopilot","Learning Role":"None",
        "Functional Scope":"Hover, translation and vertical-to-horizontal transition","Functional Scope Code":3,
        "Validation":"Outdoor free flight","Validation Code":5,"Hardware Evidence":"Yes","Free Flight":"Yes",
        "Key Metrics":"946 g vehicle; about 20 N maximum static thrust; indoor/outdoor hover and transition flight",
        "Project Relevance":"Near-exact propulsion and gimbal architecture with unusually detailed full-text hardware evidence.",
        "Evidence Status":"Full text verified","Evidence Note":"Peer-reviewed journal paper; earlier conference version noted by authors.",
        "DOI":"10.1177/17568293241254045","Source URL":"https://doi.org/10.1177/17568293241254045"
    },
    {
        "ID":"N02","Year":2024,"Title":"Design, Modeling, and Control of a Coaxial Drone",
        "Citation":"Chen et al.","Venue":"IEEE Transactions on Robotics","Publication Type":"Journal","Peer Reviewed":"Yes",
        "Architecture":"Two contra-rotating rotors with serial dual-axis servo rotation","Vehicle Form":"Compact elongated coaxial drone",
        "Vectoring":"Independent dual-axis rotation of the coaxial propulsion assembly","Control Family":"Nonlinear / allocation",
        "Control Method":"Six-DOF nonlinear model and nonlinear control allocation with damping injection","Learning Role":"None",
        "Functional Scope":"Position and yaw stabilization with maneuvering trajectory","Functional Scope Code":3,
        "Validation":"Hardware experiments","Validation Code":4,"Hardware Evidence":"Yes","Free Flight":"Yes",
        "Key Metrics":"Numerical simulations and physical experiments reported in IEEE T-RO",
        "Project Relevance":"One of the closest peer-reviewed architecture matches and missing from the original corpus.",
        "Evidence Status":"Full text verified","Evidence Note":"Publisher DOI and author-hosted accepted PDF located.",
        "DOI":"10.1109/TRO.2024.3354161","Source URL":"https://doi.org/10.1109/TRO.2024.3354161"
    },
    {
        "ID":"N03","Year":2026,"Title":"Modeling, Control Stabilization and Parameter Identification of a Thrust Vectoring Rocket-Type Aerial Robot",
        "Citation":"Chih and Peng","Venue":"Applied Mathematical Modelling","Publication Type":"Journal","Peer Reviewed":"Yes",
        "Architecture":"Gimbal-based coaxial rotor system with two servos","Vehicle Form":"Rocket-type coaxial aerial robot",
        "Vectoring":"Two-axis servo gimbal; differential coaxial torque","Control Family":"Identification / model-based",
        "Control Method":"Filtering-operator least squares with PSO tuning; PD stabilization","Learning Role":"Parameter identification",
        "Functional Scope":"Attitude and altitude stabilization for identification","Functional Scope Code":1,
        "Validation":"Simulation","Validation Code":1,"Hardware Evidence":"No","Free Flight":"No",
        "Key Metrics":"Noisy-measurement nonlinear parameter identification demonstrated numerically",
        "Project Relevance":"Exact architecture; useful for model identification, but it is simulation evidence only.",
        "Evidence Status":"Publisher full text/record verified","Evidence Note":"Volume 158, article 117000; published online April 2026.",
        "DOI":"10.1016/j.apm.2026.117000","Source URL":"https://doi.org/10.1016/j.apm.2026.117000"
    },
    {
        "ID":"N04","Year":2026,"Title":"Nonlinear Modeling and Energy-Based Flight Control of a Coaxial VTOL UAV with Independent Thrust Vectoring for Autonomous Landing Maneuvers",
        "Citation":"Durán-Delfín et al.","Venue":"Drones","Publication Type":"Journal","Peer Reviewed":"Yes",
        "Architecture":"Coaxial VTOL with independently tilting propulsion units","Vehicle Form":"Coaxial thrust-vectoring VTOL UAV",
        "Vectoring":"Independent propulsion tilt with nonlinear allocation","Control Family":"Nonlinear / passivity-based",
        "Control Method":"Quaternion IDA-PBC and nonlinear control allocation","Learning Role":"None",
        "Functional Scope":"Trajectory tracking and autonomous landing","Functional Scope Code":3,
        "Validation":"Simulation","Validation Code":1,"Hardware Evidence":"No","Free Flight":"No",
        "Key Metrics":"Three-dimensional numerical hover, cruise, transition and landing studies",
        "Project Relevance":"Close coaxial TVC comparator; propulsion topology is not the same single gimballed unit.",
        "Evidence Status":"Full text verified","Evidence Note":"Peer-reviewed open-access journal article published July 2026.",
        "DOI":"10.3390/drones10070512","Source URL":"https://doi.org/10.3390/drones10070512"
    },
]
main.extend(new_records)

distance = {
    "R01":1,"R02":1,"R04":1,"R05":1,"R06":1,"N01":1,"N02":1,"N03":1,
    "N04":2,"R09":2,"R20":2,"R10":3,"R13":3,
}
scope_class = {1:"Attitude",2:"Pose",3:"Trajectory",4:"Online guidance",5:"Integrated mission"}
for r in main:
    r["Architecture Distance Code"] = distance[r["ID"]]
    r["Architecture Class"] = {1:"Near-exact single/coaxial 2-axis TVC",2:"Limited-axis or alternate coaxial TVC",3:"Distributed multi-rotor vectoring"}[distance[r["ID"]]]
    r["Scope Class"] = scope_class[r["Functional Scope Code"]]

main.sort(key=lambda r: (r["Architecture Distance Code"], -r["Validation Code"], r["Year"], r["ID"]))

inspiration = []
for pid, rationale in [
    ("R32", "Neural-Fly: offline learned residual basis plus online adaptive coefficients; useful for disturbance compensation, not a direct TVC competitor."),
    ("R39", "Residual RL: adds a learned correction to a fixed controller; useful as a controller architecture, not evidence on a TVC aircraft."),
]:
    r = compact_existing(by_id[pid])
    r["Why outside main plot"] = rationale
    inspiration.append(r)

def exclusion_reason(r):
    title = str(r.get("Paper Title") or "").lower()
    pub = str(r.get("Publication Type") or "").lower()
    peer = str(r.get("Peer Reviewed") or "").lower()
    if "review" in pub or "survey" in pub or "review" in title or "survey" in title:
        return "Review/survey; not a primary experimental or modeling study."
    if "preprint" in pub or "unconfirmed" in peer or peer.startswith("no"):
        return "Standalone preprint or publication status not confirmed."
    if r["Relevance Ring"] in ("3 Analogous", "4 Method"):
        return "Physical architecture/control problem is too distant for the main plot."
    return "Peer-reviewed but architecture is farther than the retained comparator set."

excluded = []
main_orig_ids = set(selected_ids)
for r in original:
    if r["ID"] in main_orig_ids or r["ID"] in {"R32","R39"}:
        continue
    excluded.append({
        "ID":r["ID"],"Title":r["Paper Title"],"Year":r["Year"],"Publication Type":r["Publication Type"],
        "Relevance Ring":r["Relevance Ring"],"Reason":exclusion_reason(r),"Source URL":r["Primary Source URL"]
    })

payload = {
    "scope": {
        "title":"Curated TVC Ashby Corpus",
        "main_count":len(main),"inspiration_count":len(inspiration),"excluded_count":len(excluded),
        "rules":[
            "Peer-reviewed primary research only in the main plot.",
            "Physical thrust vectoring must be central to the vehicle, not merely a transferable method.",
            "Near-exact coaxial/single-propulsion two-axis systems are retained regardless of evidence level.",
            "Only a small comparator set of alternate servo-vectoring architectures is retained.",
            "Reviews, standalone preprints, generic vehicles and method-only papers are excluded from the main plot.",
            "Neural-Fly and residual RL are retained separately as method inspiration."
        ]
    },
    "main":main,"inspiration":inspiration,"excluded":excluded,
}
(OUT / "curated_corpus.json").write_text(json.dumps(payload, indent=2, ensure_ascii=False), encoding="utf-8")

# Focused Ashby plot.
W,H=2000,1250
img=Image.new("RGB",(W,H),"#f7f5ef")
draw=ImageDraw.Draw(img)
font_path=r"C:\Windows\Fonts\arial.ttf"; bold_path=r"C:\Windows\Fonts\arialbd.ttf"
font=ImageFont.truetype(font_path,20); small=ImageFont.truetype(font_path,16); tiny=ImageFont.truetype(font_path,14)
title=ImageFont.truetype(bold_path,34); panel=ImageFont.truetype(bold_path,22)
draw.text((55,35),"Scoped Ashby Map: Comparable Thrust-Vector Vehicles",fill="#202020",font=title)
draw.text((55,85),f"{len(main)} peer-reviewed primary studies. Method-only papers, reviews and standalone preprints are not plotted.",fill="#555555",font=font)
left,top,right,bottom=220,180,1910,940
draw.rectangle((left,top,right,bottom),fill="#fffdfa")
plot_l,plot_t,plot_r,plot_b=370,245,1840,810
validation={0:"Review",1:"Simulation",2:"SIL / HIL",3:"Constrained rig",4:"Indoor / hardware flight",5:"Outdoor free flight",6:"Operational"}
xlabels={1:"Near-exact\nsingle/coaxial\n2-axis TVC",2:"Limited-axis or\nalternate coaxial\nTVC",3:"Distributed\nmulti-rotor\nvectoring"}
for y in range(1,6):
    py=plot_b-(y-1)/4*(plot_b-plot_t)
    draw.line((plot_l,py,plot_r,py),fill="#ddd9d0",width=1)
    lab=validation[y]; tw=draw.textbbox((0,0),lab,font=small)[2]
    draw.text((plot_l-tw-16,py-10),lab,fill="#333",font=small)
for x in range(1,4):
    px=plot_l+(x-1)/2*(plot_r-plot_l)
    draw.line((px,plot_t,px,plot_b),fill="#e6e2d8",width=1)
    bb=draw.multiline_textbbox((0,0),xlabels[x],font=small,align="center")
    draw.multiline_text((px-(bb[2]-bb[0])/2,plot_b+18),xlabels[x],fill="#333",font=small,align="center",spacing=3)
draw.line((plot_l,plot_t,plot_l,plot_b),fill="#333",width=2); draw.line((plot_l,plot_b,plot_r,plot_b),fill="#333",width=2)
draw.text((left+20,top+18),"Architecture distance vs. validation maturity",fill="#202020",font=panel)

colors={"Learning / RL":"#d95f02","Optimization / MPC":"#1b9e77","Robust / nonlinear":"#7570b3","Identification":"#666666","Classical / allocation":"#1f78b4"}
def group(name):
    s=name.lower()
    if "reinforcement" in s or "learning" in s: return "Learning / RL"
    if "optimization" in s or "mpc" in s: return "Optimization / MPC"
    if "robust" in s or "nonlinear" in s or "passivity" in s: return "Robust / nonlinear"
    if "identification" in s: return "Identification"
    return "Classical / allocation"
shapes={1:"o",2:"s",3:"^",4:"D",5:"D"}
def marker(cx,cy,shape,fill,r=10):
    if shape=="o": draw.ellipse((cx-r,cy-r,cx+r,cy+r),fill=fill,outline="#202020",width=2)
    elif shape=="s": draw.rectangle((cx-r,cy-r,cx+r,cy+r),fill=fill,outline="#202020",width=2)
    elif shape=="^": draw.polygon([(cx,cy-r-2),(cx-r-2,cy+r),(cx+r+2,cy+r)],fill=fill,outline="#202020")
    else: draw.polygon([(cx,cy-r-2),(cx-r-2,cy),(cx,cy+r+2),(cx+r+2,cy)],fill=fill,outline="#202020")
rng=random.Random(42)
for r in main:
    x=r["Architecture Distance Code"]; y=r["Validation Code"]
    base_x=plot_l+(x-1)/2*(plot_r-plot_l)
    if x == 1:
        px=base_x+rng.uniform(-25,210)
    elif x == 3:
        px=base_x+rng.uniform(-210,25)
    else:
        px=base_x+rng.uniform(-130,130)
    py=plot_b-(y-1)/4*(plot_b-plot_t)+rng.uniform(-22,22)
    marker(px,py,shapes[r["Functional Scope Code"]],colors[group(r["Control Family"])])
    draw.text((px+13,py-18),r["ID"],fill="#202020",font=tiny)

draw.rectangle((plot_l+520,plot_t+18,plot_l+800,plot_t+82),fill="#fff1c7",outline="#d9b44a",width=1)
draw.multiline_text((plot_l+545,plot_t+27),"Sparse target: advanced control +\nfree-flight evidence on the exact platform",fill="#6b4f00",font=tiny,align="center",spacing=2)

ly=1010
draw.text((70,ly),"Control family (color)",fill="#202020",font=ImageFont.truetype(bold_path,16))
x=70
for lab,col in colors.items():
    marker(x+8,ly+40,"o",col,7); draw.text((x+22,ly+29),lab,fill="#333",font=tiny); x+=290
draw.text((70,ly+78),"Functional scope (shape)",fill="#202020",font=ImageFont.truetype(bold_path,16))
x=70
for code,lab in [(1,"Attitude"),(2,"Pose"),(3,"Trajectory"),(4,"Guidance/mission")]:
    marker(x+8,ly+119,shapes[code],"#dddddd",7); draw.text((x+22,ly+108),lab,fill="#333",font=tiny); x+=230
draw.text((1120,ly+108),"Lower architecture-distance code means greater physical comparability.",fill="#555",font=small)

plot_path=OUT/"tvc_ashby_scoped.png"; img.save(plot_path,dpi=(180,180))

# Reader markdown.
lines=["# Curated TVC Ashby research reader","",f"Main Ashby corpus: **{len(main)} peer-reviewed primary studies**.","",
       "![Scoped TVC Ashby plot](tvc_ashby_scoped.png)","","## Scope rule",""]
lines += [f"- {x}" for x in payload["scope"]["rules"]]
lines += ["","## What changed","",
          f"- Main plotted corpus reduced from 45 to **{len(main)}** studies.",
          f"- **{len(excluded)}** original records moved out of scope.",
          "- Four directly relevant peer-reviewed papers missing from the original corpus were added as N01–N04.",
          "- Neural-Fly and residual RL are retained only as method inspiration.","","## Main plotted papers",""]
for r in main:
    lines += [f"### {r['ID']} — [{r['Title']}]({r['Source URL']})","",
              f"**Citation:** {r['Citation']} ({r['Year']}), {r['Venue']}.  ",
              f"**Architecture:** {r['Architecture']} ({r['Architecture Class']}).  ",
              f"**Vectoring:** {r['Vectoring']}.  ",
              f"**Control:** {r['Control Family']} — {r['Control Method']}.  ",
              f"**Function:** {r['Functional Scope']}.  ",
              f"**Validation:** {r['Validation']}; hardware: {r['Hardware Evidence']}; free flight: {r['Free Flight']}.  ",
              f"**Evidence:** {r['Key Metrics']}.  ",
              f"**Why it belongs:** {r['Project Relevance']}  ",
              f"**Evidence status:** {r['Evidence Status']}. {r['Evidence Note'] or ''}",""]
lines += ["## Method inspiration — not plotted","",
          "These papers can motivate controller design, but they cannot establish the experimental state of the art for your architecture.",""]
for r in inspiration:
    lines += [f"### {r['ID']} — [{r['Title']}]({r['Source URL']})","",r["Why outside main plot"],""]
lines += ["## Excluded records","","| ID | Year | Paper | Reason |","|---|---:|---|---|"]
for r in excluded:
    title=str(r['Title']).replace('|','\\|')
    lines.append(f"| {r['ID']} | {r['Year']} | [{title}]({r['Source URL']}) | {r['Reason']} |")
lines += ["","## Interpretation caution","",
          "Sparse cells describe this curated corpus, not universal absence. Before a novelty claim, search citations around the closest exact-architecture papers, especially N02, R04, R05, R06 and N01."]
(OUT/"TVC_Ashby_Scoped_Reader.md").write_text("\n".join(lines),encoding="utf-8")

print(json.dumps({"main":len(main),"inspiration":len(inspiration),"excluded":len(excluded),"output":str(OUT)},indent=2))
